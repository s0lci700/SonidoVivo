"""Convierte un documento .md de esta carpeta a .docx para subirlo a AVA.

Uso:  python documentacion/md-a-docx.py documentacion/cobertura-testing-EP2.md
Crea el .docx junto al .md. Necesita python-docx (pip install python-docx).
Entiende lo que usan estos documentos: títulos, párrafos, listas, tablas,
citas, negrita, cursiva y código en línea.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

LATON = RGBColor(0x7A, 0x5A, 0x16)
TINTA = RGBColor(0x1A, 0x17, 0x14)


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def agregar_texto(parrafo, texto, negrita_base=False, tam=None):
    """Agrega runs con **negrita**, *cursiva* y `código`."""
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*(?!\s).+?(?<!\s)\*(?!\w))", texto)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            # La negrita puede contener código: se procesa por dentro.
            for sub in re.split(r"(`[^`]+`)", tok[2:-2]):
                if not sub:
                    continue
                es_codigo = sub.startswith("`") and sub.endswith("`")
                run = parrafo.add_run(sub[1:-1] if es_codigo else sub)
                run.bold = True
                if es_codigo:
                    run.font.name = "Consolas"
                if tam:
                    run.font.size = tam
            continue
        if tok.startswith("`") and tok.endswith("`"):
            run = parrafo.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run.font.size = tam or Pt(9.5)
            run.bold = negrita_base
            continue
        if tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            run = parrafo.add_run(tok[1:-1])
            run.italic = True
            run.bold = negrita_base
            if tam:
                run.font.size = tam
            continue
        tok = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", tok)
        run = parrafo.add_run(tok)
        run.bold = negrita_base
        if tam:
            run.font.size = tam


def celdas(linea):
    linea = linea.strip().strip("|")
    partes, actual, en_codigo = [], "", False
    for c in linea:
        if c == "`":
            en_codigo = not en_codigo
        if c == "|" and not en_codigo:
            partes.append(actual)
            actual = ""
        else:
            actual += c
    partes.append(actual)
    return [p.strip() for p in partes]


def convertir(origen):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for nombre in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[nombre].font.color.rgb = LATON if nombre != "Heading 1" else TINTA

    lineas = Path(origen).read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lineas):
        l = lineas[i]
        m = re.match(r"^(#{1,4})\s+(.*)", l)
        if m:
            p = doc.add_heading(level=min(len(m.group(1)), 3))
            agregar_texto(p, m.group(2))
            i += 1
            continue
        if l.startswith("|") and i + 1 < len(lineas) and re.match(r"^\|[\s:|-]+\|?\s*$", lineas[i + 1]):
            cabecera = celdas(l)
            i += 2
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                filas.append(celdas(lineas[i]))
                i += 1
            tabla = doc.add_table(rows=1, cols=len(cabecera))
            tabla.style = "Table Grid"
            tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, c in enumerate(cabecera):
                celda = tabla.rows[0].cells[j]
                celda.text = ""
                agregar_texto(celda.paragraphs[0], c or " ", negrita_base=True, tam=Pt(9.5))
                sombrear(celda, "EDE4D0")
            for f in filas:
                fila = tabla.add_row().cells
                for j in range(len(cabecera)):
                    fila[j].text = ""
                    agregar_texto(fila[j].paragraphs[0], f[j] if j < len(f) else "", tam=Pt(9.5))
            doc.add_paragraph()
            continue
        if l.startswith(">"):
            bloque = []
            while i < len(lineas) and lineas[i].startswith(">"):
                bloque.append(lineas[i].lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14)
            agregar_texto(p, " ".join(bloque))
            for run in p.runs:
                run.font.color.rgb = LATON
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", l)
        if m:
            texto = m.group(3)
            i += 1
            while i < len(lineas) and lineas[i].startswith("  ") and lineas[i].strip() and not re.match(r"^\s*([-*]|\d+\.)\s", lineas[i]):
                texto += " " + lineas[i].strip()
                i += 1
            estilo = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            if len(m.group(1)) >= 2:
                estilo += " 2"
            p = doc.add_paragraph(style=estilo)
            agregar_texto(p, texto)
            continue
        if not l.strip() or re.match(r"^-{3,}\s*$", l):
            i += 1
            continue
        parrafo = [l.strip()]
        i += 1
        while i < len(lineas) and lineas[i].strip() and not re.match(r"^(#|\||>|```|\s*([-*]|\d+\.)\s)", lineas[i]):
            parrafo.append(lineas[i].strip())
            i += 1
        agregar_texto(doc.add_paragraph(), " ".join(parrafo))

    destino = Path(origen).with_suffix(".docx")
    doc.save(destino)
    print(f"Generado {destino.name}")


if __name__ == "__main__":
    for archivo in sys.argv[1:]:
        convertir(archivo)
